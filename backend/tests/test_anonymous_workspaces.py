from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.config import get_settings
from app.core.database import Base, get_db
from app.main import app
from app.models import KnowledgeBase, RoleProfile, User


def _docx_bytes() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading("测试候选人", level=1)
    document.add_paragraph("拥有三年产品设计与跨团队交付经验，负责需求分析、用户研究和项目推进。")
    document.add_paragraph("使用数据分析验证产品迭代效果，并持续改进用户体验。")
    document.save(buffer)
    return buffer.getvalue()


def test_anonymous_browsers_are_isolated_and_share_published_roles(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCAL_WORKSPACE_MODE", "false")
    monkeypatch.setenv("ANONYMOUS_WORKSPACE_MODE", "true")
    monkeypatch.setenv("WORKSPACE_COOKIE_SECURE", "false")
    monkeypatch.setenv("JWT_SECRET", "test-only-anonymous-workspace-secret")
    monkeypatch.setenv("MODEL_CONFIG_ENCRYPTION_KEY", "test-only-model-config-secret")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    get_settings.cache_clear()

    engine = create_engine(
        f"sqlite:///{tmp_path / 'anonymous.db'}",
        connect_args={"check_same_thread": False},
    )
    testing_session = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(engine)

    def override_db():
        db = testing_session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_db
    try:
        browser_a = TestClient(app)
        browser_b = TestClient(app)

        first = browser_a.get("/api/resumes")
        assert first.status_code == 200
        assert "HttpOnly" in first.headers["set-cookie"]
        assert "SameSite=lax" in first.headers["set-cookie"]

        uploaded = browser_a.post(
            "/api/resumes",
            files={
                "file": (
                    "candidate.docx",
                    _docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert uploaded.status_code == 201
        resume_id = uploaded.json()["id"]

        assert len(browser_a.get("/api/resumes").json()) == 1
        assert browser_b.get("/api/resumes").json() == []
        assert browser_b.get(f"/api/resumes/{resume_id}").status_code == 404

        configured = browser_a.put(
            "/api/settings/model",
            json={
                "api_key": "sk-browser-a-only-1234",
                "base_url": "https://api.example.com/v1",
                "chat_model": "example-chat",
            },
        )
        assert configured.status_code == 200
        assert configured.json()["api_key_hint"].endswith("1234")
        assert browser_b.get("/api/settings/model").json()["api_key_configured"] is False

        with testing_session() as db:
            owner = User(
                email="knowledge@jobpilot.system",
                password_hash="managed-workspace-no-login",
            )
            db.add(owner)
            db.flush()
            base = KnowledgeBase(
                user_id=owner.id,
                name="产品经理知识包",
                role_key="product_manager",
                version="v1",
                status="published",
                embedding_model="test-embedding",
                manifest={"rules": []},
            )
            db.add(base)
            db.flush()
            db.add(
                RoleProfile(
                    knowledge_base_id=base.id,
                    role_key="product_manager",
                    name="产品经理",
                )
            )
            db.commit()

        assert browser_a.get("/api/role-profiles").json()[0]["name"] == "产品经理"
        assert browser_b.get("/api/role-profiles").json()[0]["name"] == "产品经理"
    finally:
        app.dependency_overrides.clear()
        get_settings.cache_clear()
